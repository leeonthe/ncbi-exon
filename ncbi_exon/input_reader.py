"""Read gene symbols from CSV or XLSX.

Auto-detects a 'Symbol' / 'gene' / 'gene_symbol' column (case-insensitive). For
NCBI's gene_info-style sheets, optionally filters on `type_of_gene == 'protein-coding'`.

Returns symbols in input order with duplicates removed (first occurrence wins).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

SYMBOL_COL_CANDIDATES = ("symbol", "gene_symbol", "gene", "genesymbol", "hgnc_symbol")
TYPE_COL_CANDIDATES = ("type_of_gene", "gene_type", "type")
PROTEIN_CODING_TOKEN = "protein-coding"


@dataclass
class InputStats:
    """Summary returned alongside the symbol list."""

    path: str
    sheet: str | None
    total_rows: int           # data rows scanned (excludes header)
    symbol_column: str        # header used (verbatim)
    type_column: str | None   # header used for filtering, if any
    after_filter: int         # rows kept after type filter
    after_dedupe: int         # unique symbols
    dropped_empty: int


def read_gene_symbols(
    path: str | Path,
    *,
    sheet: str | None = None,
    filter_protein_coding: bool = True,
    limit: int | None = None,
) -> tuple[list[str], InputStats]:
    """Read gene symbols from a CSV or XLSX file.

    Parameters
    ----------
    path : input file path (.csv or .xlsx).
    sheet : XLSX-only; sheet name. If None, uses the first sheet.
    filter_protein_coding : if a `type_of_gene` column is present and this is
        True, keep only rows where the value is 'protein-coding' (case-insensitive).
        If the column is absent, this flag has no effect.
    limit : if set, return at most this many unique symbols (after dedupe).
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        header, rows, sheet_name = _read_xlsx(p, sheet)
    elif suffix == ".csv":
        header, rows = _read_csv(p)
        sheet_name = None
    else:
        raise ValueError(f"Unsupported input format: {suffix}")

    sym_idx = _find_column(header, SYMBOL_COL_CANDIDATES)
    if sym_idx is None:
        # Fall back to first column.
        sym_idx = 0
    type_idx = _find_column(header, TYPE_COL_CANDIDATES) if filter_protein_coding else None

    seen: set[str] = set()
    symbols: list[str] = []
    total = 0
    kept_after_filter = 0
    dropped_empty = 0
    for row in rows:
        total += 1
        if sym_idx >= len(row):
            continue
        raw = row[sym_idx]
        symbol = _normalize_symbol(raw)
        if not symbol:
            dropped_empty += 1
            continue
        if type_idx is not None and type_idx < len(row):
            tv = str(row[type_idx] or "").strip().lower()
            if tv != PROTEIN_CODING_TOKEN:
                continue
        kept_after_filter += 1
        if symbol in seen:
            continue
        seen.add(symbol)
        symbols.append(symbol)
        if limit is not None and len(symbols) >= limit:
            break

    return symbols, InputStats(
        path=str(p),
        sheet=sheet_name,
        total_rows=total,
        symbol_column=header[sym_idx] if sym_idx < len(header) else f"col{sym_idx}",
        type_column=header[type_idx] if type_idx is not None and type_idx < len(header) else None,
        after_filter=kept_after_filter,
        after_dedupe=len(symbols),
        dropped_empty=dropped_empty,
    )


# ---------------------------------------------------------------- readers

def _read_xlsx(path: Path, sheet: str | None) -> tuple[list[str], Iterator[tuple], str]:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("openpyxl is required to read .xlsx files. Install with: pip install openpyxl") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = sheet or wb.sheetnames[0]
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], iter([]), sheet_name
    header = [str(c).strip() if c is not None else "" for c in header_row]
    return header, _materialize_rows(rows_iter), sheet_name


def _materialize_rows(rows_iter: Iterable[tuple]) -> Iterator[tuple]:
    """Coerce openpyxl row tuples to plain tuples; skip fully-empty rows."""
    for row in rows_iter:
        if row is None:
            continue
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        yield row


def _read_csv(path: Path) -> tuple[list[str], list[tuple]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        try:
            header = next(reader)
        except StopIteration:
            return [], []
        header = [c.strip() for c in header]
        rows = [tuple(row) for row in reader]
    return header, rows


def _find_column(header: list[str], candidates: tuple[str, ...]) -> int | None:
    norm = [h.strip().lstrip("#").lower() for h in header]
    for cand in candidates:
        if cand in norm:
            return norm.index(cand)
    return None


def _normalize_symbol(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if not s or s in ("-", "NA", "N/A", "None"):
        return ""
    return s.upper()
